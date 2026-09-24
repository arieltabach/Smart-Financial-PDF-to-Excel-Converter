import io
import subprocess
import sys
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from converter.data_cleaner import rows_to_dataframe, validate_running_balance
from converter.excel_exporter import to_excel_bytes
from converter.pdf_extractor import _table_to_text, contains_hebrew, extract_pdf, fix_rtl_text, looks_reversed

ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture(scope="session")
def sample_pdf(tmp_path_factory) -> bytes:
    out = tmp_path_factory.mktemp("pdf") / "demo.pdf"
    subprocess.run([sys.executable, str(ROOT / "scripts" / "make_sample_pdf.py"), str(out)], check=True)
    return out.read_bytes()


def test_extract_pdf_text_and_tables(sample_pdf):
    doc = extract_pdf(sample_pdf)
    assert doc.n_pages == 1
    assert doc.n_tables >= 1
    assert "SUPER-PHARM" in doc.pages[0]
    assert "12,500.00" in doc.pages[0]
    assert doc.pages[0].startswith("[TABLES]")
    assert "[TEXT]" in doc.pages[0]
    assert "Description: Wolt delivery | Reference: 4471 | Debit: 74.50 | Credit: - | Balance: 9,296.24" in doc.pages[0]
    assert not doc.has_hebrew and not doc.rtl_fixed


def test_table_to_text_labels_cells_with_header():
    table = [["Date", "Desc", "Debit", "Credit"], ["1/1/25", "x", "5", None], [None, None, None, None]]
    out = _table_to_text(table).splitlines()
    assert out[0] == "Date | Desc | Debit | Credit"
    assert out[1] == "Date: 1/1/25 | Desc: x | Debit: 5 | Credit: -"
    assert len(out) == 2


def test_table_to_text_without_header_uses_pipes():
    out = _table_to_text([["1/1/25", "x", "5", ""], ["2/1/25", "y", "", "7"]])
    assert out == "1/1/25 | x | 5 | -\n2/1/25 | y | - | 7"


def test_contains_hebrew():
    assert contains_hebrew("יתרה 100")
    assert not contains_hebrew("balance 100")


def test_fix_rtl_text_restores_logical_order():
    logical = "שלום 100.00 עולם"
    # Visual-order PDFs mirror the line but keep digits left-to-right:
    visual = "םלוע 100.00 םולש"
    assert fix_rtl_text(visual) == logical
    assert fix_rtl_text("plain 123") == "plain 123"


def test_looks_reversed_heuristic():
    assert looks_reversed("ךיראת  הרתי  הבוח")
    assert not looks_reversed("תאריך  יתרה  חובה")


def _df():
    rows = [
        {"date": "02/02/2025", "description": "קפה", "reference": "4471", "debit": "20.00", "credit": "", "balance": "980.00"},
        {"date": "03/02/2025", "description": "משכורת", "reference": "", "debit": "", "credit": "5,000.00", "balance": "5,000.00"},  # mismatch
    ]
    return validate_running_balance(rows_to_dataframe(rows))


def test_excel_roundtrip_formatting_and_rtl():
    df, report = _df()
    data = to_excel_bytes(df, report=report, rtl=True, source_name="demo.pdf")
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Transactions", "Summary"]
    ws = wb["Transactions"]
    assert ws.sheet_view.rightToLeft is True
    assert [c.value for c in ws[1]] == ["Date", "Description", "Reference", "Debit", "Credit", "Amount", "Balance", "Balance OK"]
    assert ws["A2"].number_format == "DD/MM/YYYY"
    assert ws["D2"].value == 20.0 and "#,##0.00" in ws["D2"].number_format
    assert ws["B2"].value == "קפה"
    assert ws["H3"].value == "MISMATCH"
    assert ws["A3"].fill.fgColor.rgb.endswith("FDE2E2")
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref
    assert ws.column_dimensions["B"].width >= 8

    summary = {r[0].value: r[1].value for r in wb["Summary"].iter_rows(min_row=2)}
    assert summary["Transactions"] == 2
    assert summary["Total credit"] == 5000.0
    assert summary["Math validation"] == "REVIEW"
    assert summary["Source file"] == "demo.pdf"


def test_excel_ltr_and_empty():
    df, report = validate_running_balance(rows_to_dataframe([]))
    wb = load_workbook(io.BytesIO(to_excel_bytes(df, report=report, rtl=False)))
    assert wb["Transactions"].sheet_view.rightToLeft is False
    assert wb["Transactions"].max_row == 1
