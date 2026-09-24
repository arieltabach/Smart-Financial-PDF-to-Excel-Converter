"""Write the cleaned transactions to a formatted ``.xlsx`` with openpyxl."""
from __future__ import annotations

import io
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .data_cleaner import COLUMNS, ValidationReport

HEADER_FILL = PatternFill("solid", fgColor="1F6F8B")
HEADER_FONT = Font(bold=True, color="FFFFFF")
MISMATCH_FILL = PatternFill("solid", fgColor="FDE2E2")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

NUMBER_FORMAT = "#,##0.00;[Red]-#,##0.00"
DATE_FORMAT = "DD/MM/YYYY"


def _autofit(ws, min_width: int = 8, max_width: int = 60) -> None:
    for col_cells in ws.columns:
        length = max((len(str(c.value)) if c.value is not None else 0) for c in col_cells)
        ws.column_dimensions[get_column_letter(col_cells[0].column)].width = max(min_width, min(max_width, length + 2))


def _style_header(ws, ncols: int) -> None:
    for col in range(1, ncols + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"


def build_workbook(
    df: pd.DataFrame,
    report: ValidationReport | None = None,
    rtl: bool = False,
    source_name: str = "",
    include_validation_column: bool = True,
) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Transactions"
    ws.sheet_view.rightToLeft = rtl

    cols = list(COLUMNS)
    if include_validation_column and "Balance OK" in df.columns:
        cols.append("Balance OK")

    ws.append(cols)
    _style_header(ws, len(cols))

    for _, row in df.iterrows():
        values = []
        for c in cols:
            v = row.get(c)
            if c == "Date":
                v = None if pd.isna(v) else pd.Timestamp(v).to_pydatetime()
            elif c in ("Debit", "Credit", "Amount", "Balance"):
                v = None if pd.isna(v) else float(v)
            elif c == "Balance OK":
                v = None if v is None or (isinstance(v, float) and pd.isna(v)) else ("OK" if v else "MISMATCH")
            values.append(v)
        ws.append(values)

    text_align = Alignment(horizontal="right" if rtl else "left", vertical="center")
    for r in range(2, ws.max_row + 1):
        mismatch = False
        for ci, c in enumerate(cols, start=1):
            cell = ws.cell(row=r, column=ci)
            cell.border = BORDER
            if c == "Date":
                cell.number_format = DATE_FORMAT
                cell.alignment = Alignment(horizontal="center")
            elif c in ("Debit", "Credit", "Amount", "Balance"):
                cell.number_format = NUMBER_FORMAT
            elif c == "Balance OK":
                cell.alignment = Alignment(horizontal="center")
                if cell.value == "MISMATCH":
                    mismatch = True
                    cell.font = Font(color="C00000", bold=True)
            else:
                cell.alignment = text_align
        if mismatch:
            for ci in range(1, len(cols) + 1):
                ws.cell(row=r, column=ci).fill = MISMATCH_FILL

    ws.auto_filter.ref = ws.dimensions
    _autofit(ws)

    # ------------------------------------------------------------- summary
    summary = wb.create_sheet("Summary")
    summary.sheet_view.rightToLeft = rtl
    n = len(df)
    lines: list[tuple[str, object]] = [
        ("Source file", source_name or "-"),
        ("Generated at", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Transactions", n),
        ("Total debit", float(df["Debit"].fillna(0).sum()) if n else 0.0),
        ("Total credit", float(df["Credit"].fillna(0).sum()) if n else 0.0),
        ("Net (credit − debit)", float(df["Amount"].fillna(0).sum()) if n else 0.0),
    ]
    if n and df["Date"].notna().any():
        lines += [
            ("First date", pd.Timestamp(df["Date"].min()).to_pydatetime()),
            ("Last date", pd.Timestamp(df["Date"].max()).to_pydatetime()),
        ]
    if report is not None:
        lines += [
            ("Balance rows checked", report.n_checked),
            ("Balance mismatches", report.n_mismatch),
            ("Statement order", report.direction),
            ("Math validation", "PASS" if report.ok else "REVIEW"),
        ]
    summary.append(["Metric", "Value"])
    _style_header(summary, 2)
    for k, v in lines:
        summary.append([k, v])
    for r in range(2, summary.max_row + 1):
        cell = summary.cell(row=r, column=2)
        if isinstance(cell.value, float):
            cell.number_format = NUMBER_FORMAT
        elif isinstance(cell.value, datetime):
            cell.number_format = DATE_FORMAT
        summary.cell(row=r, column=1).font = Font(bold=True)
    _autofit(summary)
    return wb


def to_excel_bytes(df: pd.DataFrame, report: ValidationReport | None = None, rtl: bool = False, source_name: str = "") -> bytes:
    wb = build_workbook(df, report=report, rtl=rtl, source_name=source_name)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
